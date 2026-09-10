"""
generate_pptx_deck.py
=====================
Builds the official 19-slide widescreen (16:9) Microsoft PowerPoint presentation
for AIRA: AI-Based Student Mental Health Monitoring and Support System.
Embeds all 6 high-resolution 300 DPI benchmark graphs and formats slides
with clean, professional styling.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    DARK_BG = RGBColor(15, 23, 42)      # #0f172a Deep Navy
    LIGHT_BG = RGBColor(248, 250, 252)  # #f8fafc Off-White
    PRIMARY = RGBColor(2, 132, 199)     # #0284c7 Sky Blue
    ACCENT = RGBColor(114, 9, 183)      # #7209b7 Purple
    TEXT_DARK = RGBColor(30, 41, 59)    # #1e293b Slate 800
    TEXT_MUTED = RGBColor(100, 116, 139)# #64748b Slate 500
    CARD_BG = RGBColor(255, 255, 255)   # Pure White
    CARD_BORDER = RGBColor(226, 232, 240) # Slate 200
    SUCCESS = RGBColor(16, 185, 129)    # Emerald Green
    DANGER = RGBColor(239, 68, 68)      # Coral Red

    graph_dir = os.path.join(os.path.dirname(__file__), "presentation_graphs")

    def add_header(slide, title_text, category="AIRA • PRACTICE SCHOOL-I DEFENSE"):
        # Category / Kicker
        tb_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.3))
        tf_cat = tb_cat.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = PRIMARY

        # Main Slide Title
        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.7), Inches(0.7))
        tf_title = tb_title.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_DARK

    def create_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
        return shape

    # =========================================================================
    # SLIDE 1: TITLE SLIDE (DARK MODERN COVER)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = DARK_BG
    bg1.line.color.rgb = DARK_BG

    # Kicker
    tb = s1.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(10.9), Inches(0.5))
    p = tb.text_frame.paragraphs[0]
    p.text = "PRACTICE SCHOOL – I FINAL EXAMINATION & TECHNICAL DEFENSE"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = RGBColor(56, 189, 248) # Neon Cyan

    # Main Title
    tb = s1.shapes.add_textbox(Inches(1.2), Inches(1.7), Inches(10.9), Inches(1.5))
    p = tb.text_frame.paragraphs[0]
    p.text = "AIRA: AI-Based Student Mental Health\nMonitoring & Support Platform"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    # Subtitle
    tb = s1.shapes.add_textbox(Inches(1.2), Inches(3.4), Inches(10.9), Inches(0.8))
    p = tb.text_frame.paragraphs[0]
    p.text = "A Decoupled Full-Stack Architecture Integrating Dual ML Risk Diagnostics, Conversational AI Crisis Triaging, and Live Spatial Specialist Referrals"
    p.font.size = Pt(15)
    p.font.color.rgb = RGBColor(203, 213, 225)

    # Details Cards
    card_stu = create_card(s1, Inches(1.2), Inches(4.5), Inches(5.2), Inches(2.2), RGBColor(30, 41, 59), PRIMARY)
    tb_stu = s1.shapes.add_textbox(Inches(1.4), Inches(4.7), Inches(4.8), Inches(1.8))
    tf_stu = tb_stu.text_frame
    p = tf_stu.paragraphs[0]
    p.text = "DEVELOPED & PRESENTED BY:"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(56, 189, 248)
    
    p2 = tf_stu.add_paragraph()
    p2.text = "• Anand Singh Rathore (Roll No: 2024BTECH158)\n• Diksha Shekhawat (Roll No: 2024BTECH156)\nDegree: B.Tech Computer Science & Engineering"
    p2.font.size = Pt(13)
    p2.font.color.rgb = RGBColor(255, 255, 255)

    card_sup = create_card(s1, Inches(6.8), Inches(4.5), Inches(5.3), Inches(2.2), RGBColor(30, 41, 59), ACCENT)
    tb_sup = s1.shapes.add_textbox(Inches(7.0), Inches(4.7), Inches(4.9), Inches(1.8))
    tf_sup = tb_sup.text_frame
    p = tf_sup.paragraphs[0]
    p.text = "SUPERVISORS & INSTITUTION:"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(192, 132, 252)

    p2 = tf_sup.add_paragraph()
    p2.text = "• Internal: Dr. Sonali Vyas & Dr. Rajnish Kumar\n• External: Dr. Saurabh Kumar\nDepartment of CSE, JK Lakshmipat University, Jaipur"
    p2.font.size = Pt(13)
    p2.font.color.rgb = RGBColor(255, 255, 255)

    # =========================================================================
    # SLIDE 2: THE PROBLEM STATEMENT & CLINICAL MOTIVATION
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "The Mental Health Crisis in Higher Education")

    # 3 Stat Cards
    c1 = create_card(s2, Inches(0.8), Inches(1.5), Inches(3.6), Inches(2.2))
    t1 = s2.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(3.2), Inches(1.8))
    t1.text_frame.paragraphs[0].text = "60%+"
    t1.text_frame.paragraphs[0].font.size = Pt(36)
    t1.text_frame.paragraphs[0].font.bold = True
    t1.text_frame.paragraphs[0].font.color.rgb = DANGER
    p = t1.text_frame.add_paragraph()
    p.text = "Of university students report acute academic anxiety, sleep deprivation, or emotional burnout."
    p.font.size = Pt(12)

    c2 = create_card(s2, Inches(4.8), Inches(1.5), Inches(3.6), Inches(2.2))
    t2 = s2.shapes.add_textbox(Inches(5.0), Inches(1.7), Inches(3.2), Inches(1.8))
    t2.text_frame.paragraphs[0].text = "< 15%"
    t2.text_frame.paragraphs[0].font.size = Pt(36)
    t2.text_frame.paragraphs[0].font.bold = True
    t2.text_frame.paragraphs[0].font.color.rgb = PRIMARY
    p = t2.text_frame.add_paragraph()
    p.text = "Seek institutional counseling due to social stigma, peer judgment, and privacy fears."
    p.font.size = Pt(12)

    c3 = create_card(s2, Inches(8.8), Inches(1.5), Inches(3.6), Inches(2.2))
    t3 = s2.shapes.add_textbox(Inches(9.0), Inches(1.7), Inches(3.2), Inches(1.8))
    t3.text_frame.paragraphs[0].text = "11 PM - 4 AM"
    t3.text_frame.paragraphs[0].font.size = Pt(28)
    t3.text_frame.paragraphs[0].font.bold = True
    t3.text_frame.paragraphs[0].font.color.rgb = ACCENT
    p = t3.text_frame.add_paragraph()
    p.text = "Peak hours for psychological distress and crisis despair when university clinics are closed."
    p.font.size = Pt(12)

    # Detailed Analysis Card
    c4 = create_card(s2, Inches(0.8), Inches(4.0), Inches(11.6), Inches(2.9))
    t4 = s2.shapes.add_textbox(Inches(1.1), Inches(4.2), Inches(11.0), Inches(2.5))
    tf4 = t4.text_frame
    tf4.paragraphs[0].text = "Key Bottlenecks in Traditional Institutional Care:"
    tf4.paragraphs[0].font.size = Pt(16)
    tf4.paragraphs[0].font.bold = True
    tf4.paragraphs[0].font.color.rgb = TEXT_DARK

    bullets = [
        "Appointment Waiting Queues: Traditional campus centers take 3 to 14 days to schedule an initial session, failing during urgent crises.",
        "Static Survey Blindspot: Paper questionnaires (PHQ-9, GAD-7) capture a single static moment, missing daily oscillations between sleep and strain.",
        "Reactive Post-Crisis Model: Support only occurs after academic failure or severe breakdown, lacking early proactive behavioral detection."
    ]
    for b in bullets:
        p = tf4.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 3: THE AIRA SOLUTION OVERVIEW
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "The AIRA Solution — Proactive 24/7 Digital Wellness")

    sol_cards = [
        ("1. Dual ML Diagnostics", "Blends Hugging Face DistilBERT NLP (free-text emotions) with Scikit-Learn Logistic Regression (sleep, study, screen time) to compute an objective 0–100 Wellness Index.", PRIMARY),
        ("2. Empathetic Conversational AI", "Llama 3.3 70B on Groq LPUs delivering sub-400ms coaching with an automated CrisisHandler that immediately intercepts self-harm and surfaces 24/7 helplines.", ACCENT),
        ("3. Live Geolocation Referrals", "Captures device GPS coordinates, queries Google Places API (New), and applies the Haversine formula to sort verified local psychiatrists within a strict 100 km boundary.", SUCCESS)
    ]

    for i, (head, desc, col) in enumerate(sol_cards):
        cx = create_card(s3, Inches(0.8 + i * 4.0), Inches(1.6), Inches(3.7), Inches(5.2))
        t = s3.shapes.add_textbox(Inches(1.0 + i * 4.0), Inches(1.9), Inches(3.3), Inches(4.6))
        tf = t.text_frame
        p = tf.paragraphs[0]
        p.text = head
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_DARK

        p3 = tf.add_paragraph()
        p3.text = "\nCore Advantage:\n• Completely Anonymous\n• Real-Time Triaging (< 1s)\n• Zero Appointment Waiting"
        p3.font.size = Pt(11)
        p3.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 4: SYSTEM ARCHITECTURE & FULL-STACK DATAFLOW
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Decoupled Full-Stack Architecture & Microservices")

    layers = [
        ("CLIENT FRONTEND (Port 3000)", "Vanilla HTML5 / CSS3 Glassmorphic UI / ES6+ JS / Chart.js\nVite 6 Development & Build Server (Zero Framework Bloat)", PRIMARY, 1.5),
        ("FLASK MICROSERVICE API (Port 5000)", "Python 3.12 / Decoupled Modular Blueprints / PyJWT Stateless Tokens\nBcrypt 12 Salt Rounds Hashing / Input Validation Middleware", ACCENT, 3.2),
        ("INTELLIGENCE & AI LAYER", "DistilBERT Multi-Emotion Singleton / Logistic Regression Predictor\nGroq LPU Llama 3.3 70B LLM / Google Places (New) API", SUCCESS, 4.9)
    ]

    for title, detail, col, y in layers:
        create_card(s4, Inches(0.8), Inches(y), Inches(11.7), Inches(1.4))
        tb = s4.shapes.add_textbox(Inches(1.1), Inches(y + 0.15), Inches(11.1), Inches(1.1))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = col
        p2 = tf.add_paragraph()
        p2.text = detail
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_DARK

    # Bottom Database & Email banner
    create_card(s4, Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.6), RGBColor(30, 41, 59))
    tb_bot = s4.shapes.add_textbox(Inches(1.1), Inches(6.55), Inches(11.1), Inches(0.5))
    p = tb_bot.text_frame.paragraphs[0]
    p.text = "PERSISTENCE & EXTERNAL SERVICES: MongoDB Atlas Cluster (5-Min TTL OTP Index) • Brevo REST API v3 (HTTPS Port 443)"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    # =========================================================================
    # SLIDE 5: DATASETS USED: DIMENSIONS, SHAPES & CLEANING
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Empirical Datasets & Data Preprocessing Pipeline")

    c_left = create_card(s5, Inches(0.8), Inches(1.5), Inches(6.5), Inches(5.4))
    t_left = s5.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(6.1), Inches(5.0))
    tf_l = t_left.text_frame
    tf_l.paragraphs[0].text = "1. Student Depression Dataset (Primary)"
    tf_l.paragraphs[0].font.size = Pt(16)
    tf_l.paragraphs[0].font.bold = True
    tf_l.paragraphs[0].font.color.rgb = PRIMARY

    data_bullets = [
        "Sample Size: 27,901 authentic student records (~2.8 MB).",
        "Total Dimensions: 18 demographic, academic, and psychological variables.",
        "Core Features: Academic Pressure (1-5), CGPA, Sleep Duration, Work/Study Hours, Study Satisfaction, Financial Stress, and Suicidal History.",
        "Target Label: Depression (Binary classification: 0 = Low/Balanced, 1 = Elevated Risk).",
        "Train/Test Split: 80% Training (22,320 rows) and 20% Testing (5,581 rows)."
    ]
    for b in data_bullets:
        p = tf_l.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK

    p_sec = tf_l.add_paragraph()
    p_sec.text = "\n2. Secondary Datasets: smmh.csv (481 rows, social media usage) & Student Mental health.csv (101 rows, panic/anxiety correlations)."
    p_sec.font.size = Pt(11)
    p_sec.font.color.rgb = TEXT_MUTED

    c_right = create_card(s5, Inches(7.6), Inches(1.5), Inches(4.9), Inches(5.4))
    t_right = s5.shapes.add_textbox(Inches(7.8), Inches(1.7), Inches(4.5), Inches(5.0))
    tf_r = t_right.text_frame
    tf_r.paragraphs[0].text = "Data Engineering & Pipeline:"
    tf_r.paragraphs[0].font.size = Pt(16)
    tf_r.paragraphs[0].font.bold = True
    tf_r.paragraphs[0].font.color.rgb = ACCENT

    prep_steps = [
        "Cleaning: Dropped corrupt entries, imputed missing numerical fields with medians and categorical fields with modes.",
        "OneHotEncoding: Categorical variables (Gender, Sleep Duration, Diet) transformed into numerical sparse vectors.",
        "StandardScaler: All continuous metrics standardized to zero mean and unit variance ($z = (x - \\mu) / \\sigma$).",
        "NLP Corpora: depression-reddit-cleaned (7,731 posts) and mental-health-corpus (27,977 texts) for linguistic tuning."
    ]
    for s in prep_steps:
        p = tf_r.add_paragraph()
        p.text = "• " + s
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 6: MODEL PERFORMANCE BENCHMARK (GRAPH 1 EMBEDDED)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Behavioral Predictor — Multi-Model Benchmark Comparison")

    img1_path = os.path.join(graph_dir, "1_model_performance_comparison.png")
    if os.path.exists(img1_path):
        s6.shapes.add_picture(img1_path, Inches(0.8), Inches(1.5), width=Inches(7.2))

    c_info = create_card(s6, Inches(8.3), Inches(1.5), Inches(4.2), Inches(5.4))
    t_info = s6.shapes.add_textbox(Inches(8.5), Inches(1.7), Inches(3.8), Inches(5.0))
    tf_info = t_info.text_frame
    tf_info.paragraphs[0].text = "Why Logistic Regression Won:"
    tf_info.paragraphs[0].font.size = Pt(16)
    tf_info.paragraphs[0].font.bold = True
    tf_info.paragraphs[0].font.color.rgb = PRIMARY

    points = [
        "Accuracy: 79.82% vs XGBoost (78.73%) and Random Forest (77.39%).",
        "Recall: 85.16% — The highest sensitivity in detecting at-risk students.",
        "ROC-AUC: 0.8666 — Dominant discriminatory power across all probability thresholds.",
        "Latency: Sub-0.5ms inference execution, requiring zero heavy dependencies.",
        "Clinical Calibrations: Outputs direct probabilities via the Sigmoid curve that directly map to our 0–100 Wellness Index."
    ]
    for pt in points:
        p = tf_info.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 7: CONFUSION MATRICES & CLINICAL RECALL (GRAPH 2 EMBEDDED)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Confusion Matrix Analysis & Clinical Safety Defense")

    img2_path = os.path.join(graph_dir, "2_confusion_matrices_comparison.png")
    if os.path.exists(img2_path):
        s7.shapes.add_picture(img2_path, Inches(0.8), Inches(1.5), width=Inches(11.7))

    c_note = create_card(s7, Inches(0.8), Inches(5.7), Inches(11.7), Inches(1.4), RGBColor(241, 245, 249))
    t_note = s7.shapes.add_textbox(Inches(1.1), Inches(5.8), Inches(11.1), Inches(1.2))
    tf_n = t_note.text_frame
    tf_n.paragraphs[0].text = "The Clinical Recall Principle (Saving Lives):"
    tf_n.paragraphs[0].font.size = Pt(14)
    tf_n.paragraphs[0].font.bold = True
    tf_n.paragraphs[0].font.color.rgb = DANGER

    p = tf_n.add_paragraph()
    p.text = "In healthcare diagnostics, a False Negative (FN) means an acutely distressed student is missed and receives no support. Logistic Regression achieved only 485 False Negatives, outperforming Random Forest (604 missed) and XGBoost (536 missed). Minimizing False Negatives is our primary clinical safety constraint."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 8: ROC-AUC CURVES (GRAPH 3 EMBEDDED)
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Receiver Operating Characteristic (ROC-AUC) Analysis")

    img3_path = os.path.join(graph_dir, "3_roc_auc_curves.png")
    if os.path.exists(img3_path):
        s8.shapes.add_picture(img3_path, Inches(0.8), Inches(1.5), width=Inches(6.8))

    c_roc = create_card(s8, Inches(7.9), Inches(1.5), Inches(4.6), Inches(5.4))
    t_roc = s8.shapes.add_textbox(Inches(8.1), Inches(1.7), Inches(4.2), Inches(5.0))
    tf_roc = t_roc.text_frame
    tf_roc.paragraphs[0].text = "Discriminative Thresholds:"
    tf_roc.paragraphs[0].font.size = Pt(16)
    tf_roc.paragraphs[0].font.bold = True
    tf_roc.paragraphs[0].font.color.rgb = PRIMARY

    roc_pts = [
        "Area Under Curve (AUC) measures model performance across all possible classification cutoffs.",
        "Logistic Regression (AUC = 0.8666) demonstrates superior separation between healthy and distressed students.",
        "Outperforms XGBoost (AUC = 0.8537) and Random Forest (AUC = 0.8444).",
        "Significantly outperforms the diagonal random guess baseline (AUC = 0.5000), proving high clinical reliability."
    ]
    for rp in roc_pts:
        p = tf_roc.add_paragraph()
        p.text = "• " + rp
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 9: BEHAVIORAL FEATURE CORRELATIONS (GRAPH 4 EMBEDDED)
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Behavioral Drivers of Student Depression (N = 27,901)")

    img4_path = os.path.join(graph_dir, "4_feature_importance_correlation.png")
    if os.path.exists(img4_path):
        s9.shapes.add_picture(img4_path, Inches(0.8), Inches(1.5), width=Inches(7.2))

    c_corr = create_card(s9, Inches(8.3), Inches(1.5), Inches(4.2), Inches(5.4))
    t_corr = s9.shapes.add_textbox(Inches(8.5), Inches(1.7), Inches(3.8), Inches(5.0))
    tf_c = t_corr.text_frame
    tf_c.paragraphs[0].text = "Key Behavioral Insights:"
    tf_c.paragraphs[0].font.size = Pt(16)
    tf_c.paragraphs[0].font.bold = True
    tf_c.paragraphs[0].font.color.rgb = ACCENT

    corr_bullets = [
        "Academic Pressure (r = +0.435): The single largest positive driver of student depression.",
        "Financial Stress (r = +0.312): Strongly amplifies academic anxiety, especially near tuition deadlines.",
        "Work/Study Hours (r = +0.285): Excess study without restorative sleep compounds cognitive fatigue.",
        "Study Satisfaction (r = -0.320): High engagement acts as the strongest protective psychological buffer."
    ]
    for cb in corr_bullets:
        p = tf_c.add_paragraph()
        p.text = "• " + cb
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 10: DISTILBERT NLP EMOTION ENGINE (GRAPH 5 EMBEDDED)
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "DistilBERT NLP Deep Learning Emotion Engine")

    img5_path = os.path.join(graph_dir, "5_distilbert_emotion_analysis.png")
    if os.path.exists(img5_path):
        s10.shapes.add_picture(img5_path, Inches(0.8), Inches(1.5), width=Inches(8.0))

    c_nlp = create_card(s10, Inches(9.1), Inches(1.5), Inches(3.4), Inches(5.4))
    t_nlp = s10.shapes.add_textbox(Inches(9.3), Inches(1.7), Inches(3.0), Inches(5.0))
    tf_nlp = t_nlp.text_frame
    tf_nlp.paragraphs[0].text = "NLP Architecture:"
    tf_nlp.paragraphs[0].font.size = Pt(16)
    tf_nlp.paragraphs[0].font.bold = True
    tf_nlp.paragraphs[0].font.color.rgb = PRIMARY

    nlp_pts = [
        "Model: distilbert-base-uncased-emotion.",
        "6 Emotional Classes: Sadness, Fear, Anger, Joy, Love, Surprise.",
        "Singleton Design: Loaded once into shared RAM (~500MB) to conserve server memory.",
        "Gibberish Filter: Detects keyboard smashes and invalid word ratios before tokenization."
    ]
    for np in nlp_pts:
        p = tf_nlp.add_paragraph()
        p.text = "• " + np
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 11: AI CHATBOT & CRISISHANDLER PROTOCOL
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "Conversational AI & CrisisHandler Safety Protocol")

    # Left: Chatbot Engine
    c_chat = create_card(s11, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.4))
    t_chat = s11.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(5.0), Inches(5.0))
    tf_ch = t_chat.text_frame
    tf_ch.paragraphs[0].text = "Groq LPU + Llama 3.3 70B Versatile"
    tf_ch.paragraphs[0].font.size = Pt(16)
    tf_ch.paragraphs[0].font.bold = True
    tf_ch.paragraphs[0].font.color.rgb = PRIMARY

    ch_pts = [
        "Sub-400ms Latency: Replaced slow 15-30s local Ollama execution with Groq's Language Processing Units.",
        "Contextual Memory (MemoryManager): Retains multi-turn session dialogue per authenticated student ID.",
        "WellnessCoach Loop: Classifies intent into venting, advice-seeking, and goal-planning modes.",
        "Strict Word Budget: Prompts capped to 160 tokens to deliver compact, actionable, empathetic answers."
    ]
    for cp in ch_pts:
        p = tf_ch.add_paragraph()
        p.text = "• " + cp
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK

    # Right: Crisis Safety
    c_safe = create_card(s11, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.4), RGBColor(254, 242, 242), DANGER)
    t_safe = s11.shapes.add_textbox(Inches(7.1), Inches(1.7), Inches(5.1), Inches(5.0))
    tf_sf = t_safe.text_frame
    tf_sf.paragraphs[0].text = "The CrisisHandler Safety Net"
    tf_sf.paragraphs[0].font.size = Pt(16)
    tf_sf.paragraphs[0].font.bold = True
    tf_sf.paragraphs[0].font.color.rgb = DANGER

    sf_pts = [
        "Pre-LLM Interception: Evaluates multi-tier regex patterns before the message ever reaches the AI.",
        "Immediate Generative Abort: Completely bypasses the LLM if self-harm or suicidal ideation is detected to eliminate hallucinations.",
        "Certified National Helplines Rendered:",
        "  - Tele-MANAS: 14416 / 1800-891-4416 (24/7 National)",
        "  - KIRAN: 1800-599-0019 (Govt. Rehabilitation)",
        "  - 988 Lifeline & Crisis Text Line (HOME to 741741)",
        "One-Click Direct Dial: Instant phone dialer launch on mobile."
    ]
    for sp in sf_pts:
        p = tf_sf.add_paragraph()
        p.text = "• " + sp
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 12: LIVE GEOLOCATION & HAVERSINE (GRAPH 6 EMBEDDED)
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "Live Specialist Geolocation & Spatial Haversine Engine")

    img6_path = os.path.join(graph_dir, "6_spatial_haversine_radius.png")
    if os.path.exists(img6_path):
        s12.shapes.add_picture(img6_path, Inches(0.8), Inches(1.5), width=Inches(6.8))

    c_geo = create_card(s12, Inches(7.9), Inches(1.5), Inches(4.6), Inches(5.4))
    t_geo = s12.shapes.add_textbox(Inches(8.1), Inches(1.7), Inches(4.2), Inches(5.0))
    tf_g = t_geo.text_frame
    tf_g.paragraphs[0].text = "Spatial Referral Logic:"
    tf_g.paragraphs[0].font.size = Pt(16)
    tf_g.paragraphs[0].font.bold = True
    tf_g.paragraphs[0].font.color.rgb = SUCCESS

    geo_pts = [
        "HTML5 GPS Auto-Detection: Captures student coordinates on page load via navigator.geolocation.",
        "Google Places API (New): Queries places.googleapis.com with field masks (displayName, rating, formattedAddress).",
        "Haversine Great-Circle Formula: Computes true spherical distance over Earth's curvature (R = 6,371 km).",
        "Strict <= 100 km Radius Cap: If Google Places hits rate limits (HTTP 429), our local database activates, strictly filtering out clinics > 100 km."
    ]
    for gp in geo_pts:
        p = tf_g.add_paragraph()
        p.text = "• " + gp
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 13: SECURITY, AUTHENTICATION & BREVO API
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    add_header(s13, "Enterprise Security, PyJWT & The Brevo HTTPS Driver")

    sec_cards = [
        ("Bcrypt (12 Salt Rounds)", "Hashes passwords using 4,096 Blowfish iterations ($2^{12}$). Hardware brute-force attacks take millennia.", PRIMARY),
        ("Stateless PyJWT Tokens", "Signed with HS256, 24-hour expiration. Eliminates server session memory bottlenecks and scales horizontally.", ACCENT),
        ("Brevo HTTPS REST API", "Bypasses Render/cloud SMTP port blocking (ports 25, 465, 587) by sending verification OTPs over HTTPS port 443.", SUCCESS),
        ("MongoDB TTL 5-Min Index", "db.otp_codes.create_index('created_at', expireAfterSeconds=300) auto-deletes expired OTPs at the database engine level.", DANGER)
    ]

    for i, (head, body, col) in enumerate(sec_cards):
        row = i // 2
        col_idx = i % 2
        c = create_card(s13, Inches(0.8 + col_idx * 5.9), Inches(1.6 + row * 2.7), Inches(5.6), Inches(2.4))
        tb = s13.shapes.add_textbox(Inches(1.0 + col_idx * 5.9), Inches(1.8 + row * 2.7), Inches(5.2), Inches(2.0))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = head
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = col
        p2 = tf.add_paragraph()
        p2.text = body
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 14: 30-DAY MOOD HEATMAP & DASHBOARD
    # =========================================================================
    s14 = prs.slides.add_slide(blank_layout)
    add_header(s14, "Student Analytics Dashboard & 30-Day Mood Heatmap")

    c_dash = create_card(s14, Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.4))
    t_dash = s14.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(11.1), Inches(4.8))
    tf_d = t_dash.text_frame
    tf_d.paragraphs[0].text = "Longitudinal Mental Health Tracking & Analytics:"
    tf_d.paragraphs[0].font.size = Pt(18)
    tf_d.paragraphs[0].font.bold = True
    tf_d.paragraphs[0].font.color.rgb = PRIMARY

    dash_bullets = [
        "GitHub-Style Contribution Heatmap: Renders a 30-day interactive grid coloring daily emotional check-ins from dark restless blue to vibrant balanced green.",
        "Click-to-Inspect Historic Logs: Students can click any previous day to inspect past sentiment distributions, journal snippets, and sleep parameters.",
        "Interactive Chart.js Visualizations: Dynamic radar charts comparing academic workload vs sleep hygiene and sentiment pie charts.",
        "Mindfulness & Grounding Audio Suite: Integrated breathing visualizer and progressive muscle relaxation audio clips to calm acute panic attacks."
    ]
    for db in dash_bullets:
        p = tf_d.add_paragraph()
        p.text = "• " + db
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 15: EMPIRICAL QA & 175 PYTESTS
    # =========================================================================
    s15 = prs.slides.add_slide(blank_layout)
    add_header(s15, "Empirical Quality Assurance — 175 Automated Pytests")

    c_t1 = create_card(s15, Inches(0.8), Inches(1.5), Inches(3.6), Inches(2.0), RGBColor(240, 253, 244), SUCCESS)
    t_t1 = s15.shapes.add_textbox(Inches(1.0), Inches(1.65), Inches(3.2), Inches(1.7))
    t_t1.text_frame.paragraphs[0].text = "175 / 175"
    t_t1.text_frame.paragraphs[0].font.size = Pt(36)
    t_t1.text_frame.paragraphs[0].font.bold = True
    t_t1.text_frame.paragraphs[0].font.color.rgb = SUCCESS
    p = t_t1.text_frame.add_paragraph()
    p.text = "Tests Passed (100% Success Rate)"
    p.font.size = Pt(13)
    p.font.bold = True

    c_t2 = create_card(s15, Inches(4.8), Inches(1.5), Inches(3.6), Inches(2.0))
    t_t2 = s15.shapes.add_textbox(Inches(5.0), Inches(1.65), Inches(3.2), Inches(1.7))
    t_t2.text_frame.paragraphs[0].text = "20.67s"
    t_t2.text_frame.paragraphs[0].font.size = Pt(36)
    t_t2.text_frame.paragraphs[0].font.bold = True
    t_t2.text_frame.paragraphs[0].font.color.rgb = PRIMARY
    p = t_t2.text_frame.add_paragraph()
    p.text = "Total Test Suite Execution Time"
    p.font.size = Pt(13)

    c_t3 = create_card(s15, Inches(8.8), Inches(1.5), Inches(3.6), Inches(2.0))
    t_t3 = s15.shapes.add_textbox(Inches(9.0), Inches(1.65), Inches(3.2), Inches(1.7))
    t_t3.text_frame.paragraphs[0].text = "0"
    t_t3.text_frame.paragraphs[0].font.size = Pt(36)
    t_t3.text_frame.paragraphs[0].font.bold = True
    t_t3.text_frame.paragraphs[0].font.color.rgb = DANGER
    p = t_t3.text_frame.add_paragraph()
    p.text = "Regression Errors or Failures"
    p.font.size = Pt(13)

    # Coverage Details
    c_cov = create_card(s15, Inches(0.8), Inches(3.8), Inches(11.6), Inches(3.1))
    t_cov = s15.shapes.add_textbox(Inches(1.1), Inches(4.0), Inches(11.0), Inches(2.7))
    tf_cov = t_cov.text_frame
    tf_cov.paragraphs[0].text = "Automated Test Coverage Scope (backend/tests/test_pytest_unit.py):"
    tf_cov.paragraphs[0].font.size = Pt(15)
    tf_cov.paragraphs[0].font.bold = True
    tf_cov.paragraphs[0].font.color.rgb = TEXT_DARK

    t_cases = [
        "Authentication & Security: Bcrypt hashing, token issuance, 24h expiration, malformed signature rejection.",
        "ML Diagnostic Pipelines: DistilBERT singleton inference, linguistic gibberish detection, Logistic Regression risk tiers.",
        "CrisisHandler Safety Net: Regex pattern matching, immediate LLM bypass, national helpline card rendering.",
        "Spatial Haversine Calculations: Distance accuracy, coordinate validation, strict <= 100 km boundary enforcement.",
        "Database Resilience: MongoDB collection indexing constraints, TTL OTP self-deletion, and mongomock in-memory fallbacks."
    ]
    for tc in t_cases:
        p = tf_cov.add_paragraph()
        p.text = "• " + tc
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 16: KEY ENGINEERING CHALLENGES SOLVED
    # =========================================================================
    s16 = prs.slides.add_slide(blank_layout)
    add_header(s16, "Key Engineering Challenges Overcome")

    challenges = [
        ("1. Cloud SMTP Port Blocking", "Render & AWS block outbound ports 25, 465, 587. Attempting SMTP led to [Errno 101] Network unreachable.", "Engineered a custom Brevo REST API email driver sending transactional OTPs over HTTPS port 443.", PRIMARY),
        ("2. Google Places API Rate Limits", "Daily unbilled quota exhaustion (HTTP 429) caused API dropouts during high-frequency testing.", "Created a local MongoDB fallback database with Haversine filtering strictly capped to <= 100 km.", ACCENT),
        ("3. Local LLM Latency (15-30s)", "Running 70B/7B models on local laptops caused high latency, spinning fans, and degraded user experience.", "Integrated Groq LPUs running Llama 3.3 70B, reducing response latency to sub-400 milliseconds.", SUCCESS)
    ]

    for i, (title, prob, sol, col) in enumerate(challenges):
        c = create_card(s16, Inches(0.8), Inches(1.5 + i * 1.8), Inches(11.7), Inches(1.6))
        tb = s16.shapes.add_textbox(Inches(1.1), Inches(1.6 + i * 1.8), Inches(11.1), Inches(1.4))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = "Problem: " + prob
        p2.font.size = Pt(12)
        p2.font.color.rgb = DANGER

        p3 = tf.add_paragraph()
        p3.text = "Solution: " + sol
        p3.font.size = Pt(12)
        p3.font.color.rgb = SUCCESS

    # =========================================================================
    # SLIDE 17: LIVE DEMONSTRATION WALKTHROUGH PLAN
    # =========================================================================
    s17 = prs.slides.add_slide(blank_layout)
    add_header(s17, "Step-by-Step Live Demonstration Protocol")

    demo_steps = [
        ("1. Launch & UI Walkthrough", "Show Vite server on http://localhost:3000/ and Flask on http://127.0.0.1:5000/. Tour the glassmorphic dark theme."),
        ("2. Signup & Password Checklist", "Show real-time visual checklist validation enforcing 8+ characters, uppercase, lowercase, and numbers."),
        ("3. Mental Health Scanner", "Enter journal: 'Exam pressure is overwhelming, haven't slept'. Show sub-second dual ML risk scoring and coping tips."),
        ("4. Chatbot Coaching Demo", "Ask for study grounding tips; demonstrate sub-400ms Llama 3.3 70B response formatting."),
        ("5. The Killer Crisis Demo", "Type: 'I feel hopeless and want to end it all'. Show immediate CrisisHandler bypass and Tele-MANAS (14416) escalation."),
        ("6. Doctor Geolocation", "Demonstrate HTML5 GPS detection pulling verified Jaipur psychiatrists with live ratings and Haversine distance badges.")
    ]

    for i, (head, desc) in enumerate(demo_steps):
        row = i // 2
        col_idx = i % 2
        c = create_card(s17, Inches(0.8 + col_idx * 5.9), Inches(1.6 + row * 1.8), Inches(5.6), Inches(1.6))
        tb = s17.shapes.add_textbox(Inches(1.0 + col_idx * 5.9), Inches(1.75 + row * 1.8), Inches(5.2), Inches(1.3))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = head
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = PRIMARY
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 18: FUTURE SCOPE & PRODUCTION ROADMAP
    # =========================================================================
    s18 = prs.slides.add_slide(blank_layout)
    add_header(s18, "Future Scope & Production Roadmap")

    c_scope = create_card(s18, Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.4))
    t_scope = s18.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(11.1), Inches(4.8))
    tf_s = t_scope.text_frame
    tf_s.paragraphs[0].text = "Strategic Next-Phase Initiatives:"
    tf_s.paragraphs[0].font.size = Pt(18)
    tf_s.paragraphs[0].font.bold = True
    tf_s.paragraphs[0].font.color.rgb = PRIMARY

    scope_pts = [
        ("Native Mobile Applications (iOS & Android)", "Developing a cross-platform React Native client with native push notifications for morning and evening mindfulness reminders."),
        ("Smartwatch HRV Sensor Integration", "Connecting Bluetooth smartwatch APIs to stream real-time Heart Rate Variability (HRV) and sleep stage data for physiological stress detection."),
        ("Encrypted WebRTC Tele-Counseling", "Enabling students to book and conduct private, end-to-end encrypted video appointments directly with licensed campus counselors."),
        ("Multi-Lingual Indic NLP Support", "Expanding DistilBERT transformer fine-tuning to support Hindi and regional Indian dialects using IndicBERT models.")
    ]
    for title, desc in scope_pts:
        p = tf_s.add_paragraph()
        p.text = "• " + title + ": " + desc
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 19: CONCLUSION & DEFENSE Q&A (DARK MODERN CLOSING)
    # =========================================================================
    s19 = prs.slides.add_slide(blank_layout)
    bg19 = s19.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg19.fill.solid()
    bg19.fill.fore_color.rgb = DARK_BG
    bg19.line.color.rgb = DARK_BG

    tb = s19.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(1.5))
    p = tb.text_frame.paragraphs[0]
    p.text = "Thank You!"
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    p2 = tb.text_frame.add_paragraph()
    p2.text = "AIRA demonstrates that artificial intelligence, behavioral analytics, and spatial services can combine to provide proactive, compassionate, and life-saving mental health care for students."
    p2.font.size = Pt(16)
    p2.font.color.rgb = RGBColor(203, 213, 225)

    card_qa = create_card(s19, Inches(1.2), Inches(4.0), Inches(10.9), Inches(2.4), RGBColor(30, 41, 59), PRIMARY)
    tb_qa = s19.shapes.add_textbox(Inches(1.5), Inches(4.2), Inches(10.3), Inches(2.0))
    tf_qa = tb_qa.text_frame
    p = tf_qa.paragraphs[0]
    p.text = "NOW OPEN FOR DEFENSE QUESTIONS & EVALUATION"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RGBColor(56, 189, 248)

    p_dt = tf_qa.add_paragraph()
    p_dt.text = "Authors: Anand Singh Rathore (2024BTECH158) & Diksha Shekhawat (2024BTECH156)\nSupervisors: Dr. Sonali Vyas, Dr. Rajnish Kumar & Dr. Saurabh Kumar\nDepartment of CSE, JK Lakshmipat University, Jaipur"
    p_dt.font.size = Pt(13)
    p_dt.font.color.rgb = RGBColor(255, 255, 255)

    # Save PowerPoint presentation
    pptx_path = os.path.join(os.path.dirname(__file__), "AIRA_Final_Presentation.pptx")
    prs.save(pptx_path)
    print(f"Successfully generated PowerPoint presentation at: {pptx_path}")

if __name__ == "__main__":
    create_presentation()
